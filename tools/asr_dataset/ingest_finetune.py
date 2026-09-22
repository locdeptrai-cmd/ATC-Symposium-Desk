"""Ingest VHF audio + Excel gold transcript: gold JSONL + learned script repairs."""
from __future__ import annotations

import csv
import json
import re
import shutil
import tempfile
import threading
import time
import unicodedata
import uuid
from datetime import datetime, time as dt_time, timedelta
from difflib import SequenceMatcher
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from asr_dataset.export_gold import build_record
from asr_dataset.paths import audio_root, ensure_gold_dirs, gold_root, manifest_path
from asr_dataset.recipe import recipe_status
from asr_dataset.vocabulary import excel_phrases, vocabulary_revision

try:
    import app_paths
except ImportError:
    app_paths = None  # type: ignore

try:
    import media_transcribe
except ImportError:
    media_transcribe = None  # type: ignore

_JOBS: dict[str, dict] = {}
_JOBS_LOCK = threading.Lock()
_SPEAKER = {"atco": "ATCO", "pilot": "PILOT", "atc": "ATCO", "twr": "ATCO", "gnd": "ATCO"}
_TIME = re.compile(r"^(\d{1,2}):(\d{2})(?::(\d{2}))?(?:\.(\d+))?$")


def learned_path() -> Path:
    ensure_gold_dirs()
    return gold_root() / "learned_repairs.json"


def glossary_phrases_path() -> Path:
    ensure_gold_dirs()
    return gold_root() / "glossary_phrases.json"


def _norm_header(value: object) -> str:
    t = str(value or "").strip().lower().replace("đ", "d")
    t = unicodedata.normalize("NFKD", t)
    t = "".join(ch for ch in t if not unicodedata.combining(ch))
    t = re.sub(r"[^a-z0-9]+", "_", t).strip("_")
    return t


def parse_clock(value: object) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, timedelta):
        return float(value.total_seconds())
    if isinstance(value, datetime):
        return value.hour * 3600 + value.minute * 60 + value.second + value.microsecond / 1_000_000
    if isinstance(value, dt_time):
        return value.hour * 3600 + value.minute * 60 + value.second + value.microsecond / 1_000_000
    if isinstance(value, (int, float)):
        return float(value)
    raw = str(value).strip()
    if not raw:
        return None
    try:
        return float(raw.replace(",", "."))
    except ValueError:
        pass
    match = _TIME.match(raw)
    if not match:
        return None
    hh, mm, ss, frac = match.group(1), match.group(2), match.group(3) or "0", match.group(4) or "0"
    extra = float("0." + frac) if frac else 0.0
    return int(hh) * 3600 + int(mm) * 60 + int(ss) + extra


def _speaker(value: object) -> str:
    t = str(value or "").strip().upper()
    if t in {"ATCO", "PILOT", "UNKNOWN"}:
        return t
    return _SPEAKER.get(t.lower(), "UNKNOWN")


def _pick(row: dict[str, object], *names: str) -> str:
    for name in names:
        if name in row and row[name] not in (None, ""):
            return str(row[name]).strip()
    return ""


def parse_excel(path: Path) -> list[dict]:
    suffix = path.suffix.lower()
    if suffix == ".xls":
        raise ValueError("File .xls không đọc được. Lưu lại thành .xlsx hoặc .csv rồi nạp lại.")
    if suffix == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as fh:
            reader = csv.reader(fh)
            rows_raw = list(reader)
        header: list[str] = []
        rows: list[dict] = []
        for raw in rows_raw:
            if not raw or not any(str(c).strip() for c in raw):
                continue
            if not header:
                header = [_norm_header(c) for c in raw]
                continue
            mapped = {header[j]: raw[j] if j < len(raw) else None for j in range(len(header))}
            rows.extend(_rows_from_mapped(mapped, len(rows)))
        return _rebase_times(rows)
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        header: list[str] = []
        rows: list[dict] = []
        for raw in ws.iter_rows(values_only=True):
            if not raw or not any(c is not None and str(c).strip() for c in raw):
                continue
            if not header:
                header = [_norm_header(c) for c in raw]
                continue
            mapped = {header[j]: raw[j] if j < len(raw) else None for j in range(len(header))}
            rows.extend(_rows_from_mapped(mapped, len(rows)))
        return _rebase_times(rows)
    finally:
        wb.close()


def _rows_from_mapped(mapped: dict[str, object], index: int) -> list[dict]:
    speaker = _speaker(_pick(mapped, "speaker", "vai", "role", "atco_pilot"))
    text = _pick(mapped, "text", "gold", "transcript", "noi_dung", "kich_ban", "script", "en")
    asr = _pick(mapped, "asr", "heard", "stt", "raw", "machine")
    t0 = parse_clock(mapped.get("t_start") or mapped.get("start") or mapped.get("bat_dau") or mapped.get("time"))
    t1 = parse_clock(mapped.get("t_end") or mapped.get("end") or mapped.get("ket_thuc"))
    atco = _pick(mapped, "atco")
    pilot = _pick(mapped, "pilot")
    out: list[dict] = []
    if atco or pilot:
        if atco:
            out.append(
                {
                    "speaker": "ATCO",
                    "text": atco,
                    "asr": asr,
                    "t_start": t0 if t0 is not None else float(index),
                    "t_end": t1 if t1 is not None else (t0 or float(index)) + 2,
                }
            )
        if pilot:
            out.append(
                {
                    "speaker": "PILOT",
                    "text": pilot,
                    "asr": asr,
                    "t_start": t1 if t1 is not None else (t0 if t0 is not None else float(index + 1)),
                    "t_end": (t1 if t1 is not None else (t0 if t0 is not None else float(index + 1))) + 2,
                }
            )
        return out
    if not text:
        return out
    script_lines = _script_lines(text)
    if len(script_lines) >= 2 and speaker == "UNKNOWN":
        return script_lines
    out.append(
        {
            "speaker": speaker if speaker != "UNKNOWN" else _guess_role(text, index),
            "text": text,
            "asr": asr,
            "t_start": t0 if t0 is not None else float(index * 4),
            "t_end": t1 if t1 is not None else float(index * 4 + 3),
        }
    )
    return out


_STRIP_STEM = ("text", "gold", "transcript", "kichban", "script", "noidung", "hoithoai")


def stem_key(name: str) -> str:
    stem = Path(str(name or "")).stem.lower().replace("đ", "d")
    stem = unicodedata.normalize("NFKD", stem)
    stem = "".join(ch for ch in stem if not unicodedata.combining(ch))
    compact = re.sub(r"[^a-z0-9]+", "", stem)
    changed = True
    while changed:
        changed = False
        for suffix in _STRIP_STEM:
            if compact.endswith(suffix) and len(compact) > len(suffix) + 2:
                compact = compact[: -len(suffix)]
                changed = True
                break
    return compact


def _rebase_times(rows: list[dict]) -> list[dict]:
    times = [float(row.get("t_start") or 0) for row in rows]
    if not times:
        return rows
    tmin = min(times)
    tmax = max(times)
    if tmin < 15 * 60 or (tmax - tmin) > 3 * 3600:
        return rows
    for row in rows:
        row["t_start"] = float(row.get("t_start") or 0) - tmin
        row["t_end"] = float(row.get("t_end") or 0) - tmin
        if row["t_end"] <= row["t_start"]:
            row["t_end"] = row["t_start"] + 2.0
    return rows


def turns_from_rows(rows: list[dict]) -> list[dict]:
    turns: list[dict] = []
    for row in rows:
        text = str(row.get("text") or "").strip()
        if not text:
            continue
        t0 = float(row.get("t_start") or 0)
        t1 = float(row.get("t_end") or 0) or t0 + 2.0
        turns.append(
            {
                "t_start": t0,
                "t_end": t1 if t1 > t0 else t0 + 2.0,
                "text": text,
                "speaker_role": str(row.get("speaker") or "UNKNOWN"),
            }
        )
    turns.sort(key=lambda item: float(item["t_start"]))
    return turns


def script_from_turns(turns: list[dict]) -> str:
    lines: list[str] = []
    for turn in turns:
        sec = max(0, int(float(turn.get("t_start") or 0)))
        mm, ss = divmod(sec, 60)
        prefix = "[%d:%02d] " % (mm, ss)
        who = str(turn.get("speaker_role") or turn.get("speaker") or "")
        if who and who != "UNKNOWN":
            prefix += who + "  "
        lines.append(prefix + str(turn.get("text") or ""))
    return "\n".join(lines)


def gold_turns_for(filename: str) -> list[dict]:
    key = stem_key(filename)
    if not key:
        return []
    path = manifest_path()
    if not path.is_file():
        return []
    matched: dict[str, list[dict]] = {}
    order: list[str] = []
    seen: dict[str, set[tuple]] = {}
    for line in path.read_text(encoding="utf-8").splitlines():
        if not line.strip():
            continue
        try:
            rec = json.loads(line)
        except json.JSONDecodeError:
            continue
        names = [str(rec.get("filename") or ""), Path(str(rec.get("audio") or "")).name]
        if not any(stem_key(item) == key for item in names if item and item != "."):
            continue
        text = str(rec.get("text") or "").strip()
        if not text:
            continue
        sid = str(rec.get("session_id") or "unknown")
        if sid not in matched:
            matched[sid] = []
            seen[sid] = set()
            order.append(sid)
        t0 = float(rec.get("t_start") or 0)
        stamp = (round(t0, 1), text.lower())
        if stamp in seen[sid]:
            continue
        seen[sid].add(stamp)
        t1 = float(rec.get("t_end") or 0) or t0 + 2.0
        matched[sid].append(
            {
                "t_start": t0,
                "t_end": t1 if t1 > t0 else t0 + 2.0,
                "text": text,
                "speaker_role": str(rec.get("speaker") or "UNKNOWN"),
            }
        )
    if not order:
        return []
    chosen = matched[order[-1]]
    chosen.sort(key=lambda item: float(item["t_start"]))
    return chosen


def gold_payload_for(filename: str) -> dict:
    turns = gold_turns_for(filename)
    return {
        "ok": True,
        "stem": stem_key(filename),
        "turns": turns,
        "text": script_from_turns(turns),
        "count": len(turns),
    }


def persist_parsed_gold(rows: list[dict], excel_name: str, audio_name: str = "") -> int:
    filename = audio_name or excel_name or "gold.xlsx"
    session_id = "reda-" + (stem_key(filename) or uuid.uuid4().hex[:8])
    return append_gold_rows(rows, filename, None, session_id)


def parse_excel_payload(path: Path, excel_name: str = "", audio_name: str = "") -> dict:
    rows = parse_excel(path)
    if not rows:
        raise ValueError("Excel không có dòng hội thoại (cần cột text/gold/speaker).")
    persist_parsed_gold(rows, excel_name or path.name, audio_name)
    turns = turns_from_rows(rows)
    return {
        "ok": True,
        "stem": stem_key(audio_name or excel_name or path.name),
        "turns": turns,
        "text": script_from_turns(turns),
        "count": len(turns),
        "filename": excel_name or path.name,
    }


def _script_lines(text: str) -> list[dict]:
    from reda.engine import parse_script

    parsed = parse_script(text, "")
    out = []
    for item in parsed:
        out.append(
            {
                "speaker": item.get("speaker_role") or "UNKNOWN",
                "text": item.get("text") or "",
                "asr": "",
                "t_start": float(item.get("t_start") or 0),
                "t_end": float(item.get("t_end") or 0) or float(item.get("t_start") or 0) + 2,
            }
        )
    return out


def _guess_role(text: str, index: int) -> str:
    t = text.upper()
    if t.startswith("ATCO"):
        return "ATCO"
    if t.startswith("PILOT"):
        return "PILOT"
    return "ATCO" if index % 2 == 0 else "PILOT"


def write_template(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "VHF"
    headers = ["t_start", "t_end", "speaker", "text", "asr"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.append(["00:00:05", "00:00:12", "ATCO", "TSN Tower VJC1225 ILSw RWY 25R", "Saigon Tower Vietjet one two two five ils whiskey runway two five right"])
    ws.append(["00:00:13", "00:00:18", "PILOT", "VJC1225 ILSw RWY 25R", "Vietjet one two two five ils whiskey runway two five right"])
    ws.append(["00:00:19", "00:00:24", "ATCO", "VJC1225 continue approach RWY 25R", "continue of course runway two seven right"])
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 48
    ws.column_dimensions["E"].width = 56
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def _phrase_sources() -> list[Path]:
    root = Path(__file__).resolve().parents[2]
    out = [
        root / "data" / "doc4444-phraseology.tsv",
        root / "data" / "vn-callsigns.tsv",
        root / "data" / "user-phraseology.json",
    ]
    if app_paths is not None:
        for guess in (
            app_paths.app_home() / "data" / "user-phraseology.json",
            app_paths.bundle_root() / "data" / "user-phraseology.json",
        ):
            if guess not in out:
                out.append(guess)
    return out


def _clean_phrase(raw: object) -> str:
    text = str(raw or "").strip()
    if not text:
        return ""
    text = re.sub(r"\([^)]*\)", " ", text)
    text = re.sub(r"[/|]+", " ", text)
    text = re.sub(r"\s+", " ", text).strip(" ,.;:-")
    words = text.split()
    if len(words) < 1 or len(words) > 8:
        return ""
    if any(ch in text for ch in "{}[]<>"):
        return ""
    if sum(ch.isalpha() for ch in text) < 3:
        return ""
    return text


def _collect_seed_phrases() -> tuple[list[str], list[str]]:
    phrases: list[str] = []
    seen: set[str] = set()
    sources: list[str] = []
    for path in _phrase_sources():
        if not path.is_file():
            continue
        sources.append(str(path))
        if path.suffix.lower() == ".json":
            try:
                payload = json.loads(path.read_text(encoding="utf-8"))
            except (OSError, json.JSONDecodeError):
                continue
            rows = payload.get("entries") if isinstance(payload, dict) else payload
            if not isinstance(rows, list):
                continue
            for row in rows:
                if not isinstance(row, dict):
                    continue
                for value in (row.get("en"), row.get("abbr")):
                    phrase = _clean_phrase(value)
                    key = phrase.lower()
                    if phrase and key not in seen:
                        seen.add(key)
                        phrases.append(phrase)
            continue
        try:
            with path.open(encoding="utf-8", newline="") as fh:
                for row in csv.DictReader(fh, delimiter="\t"):
                    phrase = _clean_phrase(row.get("en"))
                    key = phrase.lower()
                    if phrase and key not in seen:
                        seen.add(key)
                        phrases.append(phrase)
                    tele = _clean_phrase(row.get("telephony"))
                    key2 = tele.lower()
                    if tele and key2 not in seen:
                        seen.add(key2)
                        phrases.append(tele)
        except OSError:
            continue
    return phrases, sources


def load_glossary_phrases() -> list[str]:
    path = glossary_phrases_path()
    if not path.is_file():
        return []
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return []
    rows = payload.get("phrases") if isinstance(payload, dict) else payload
    if not isinstance(rows, list):
        return []
    out: list[str] = []
    seen: set[str] = set()
    for raw in rows:
        phrase = _clean_phrase(raw)
        key = phrase.lower()
        if phrase and key not in seen:
            seen.add(key)
            out.append(phrase)
    return out


def seed_glossary_phrases(limit: int = 1200) -> dict:
    phrases, sources = _collect_seed_phrases()
    payload = {
        "updated": time.time(),
        "count": min(len(phrases), max(1, limit)),
        "sources": sources,
        "phrases": phrases[: max(1, limit)],
    }
    path = glossary_phrases_path()
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return {"ok": True, "count": int(payload["count"]), "sources": sources, "path": str(path)}


def load_learned() -> list[dict]:
    path = learned_path()
    if not path.is_file():
        return []
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return []
    if isinstance(data, dict):
        data = data.get("rules") or []
    out = []
    for row in data:
        src = str(row.get("src") or "").strip()
        dst = str(row.get("dst") or "").strip()
        if src and src.lower() != dst.lower():
            out.append({"src": src, "dst": dst, "count": int(row.get("count") or 1)})
    return out


def save_learned(rules: list[dict]) -> Path:
    path = learned_path()
    payload = {"updated": time.time(), "rules": rules}
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    try:
        from reda.script_polish import reload_learned

        reload_learned()
    except Exception:
        pass
    return path


def _tokens(text: str) -> list[str]:
    return [t for t in re.findall(r"[A-Za-z0-9']+|RWY|ILSw?", text) if t]


def learn_from_pair(asr: str, gold: str, bucket: dict[tuple[str, str], int]) -> None:
    a = " ".join((asr or "").split())
    g = " ".join((gold or "").split())
    if not a or not g or a.lower() == g.lower():
        return
    at, gt = _tokens(a), _tokens(g)
    if not at or not gt:
        return
    matcher = SequenceMatcher(a=at, b=gt, autojunk=False)
    for tag, i1, i2, j1, j2 in matcher.get_opcodes():
        if tag != "replace":
            continue
        src = " ".join(at[i1:i2]).strip()
        dst = " ".join(gt[j1:j2]).strip()
        if not src or src.lower() == dst.lower():
            continue
        if len(src) < 3 or len(src.split()) > 6 or len(dst.split()) > 6:
            continue
        if src.isdigit():
            continue
        bucket[(src, dst)] = bucket.get((src, dst), 0) + 1


def merge_learned(new_pairs: dict[tuple[str, str], int]) -> list[dict]:
    current = {(r["src"], r["dst"]): int(r.get("count") or 1) for r in load_learned()}
    for key, n in new_pairs.items():
        current[key] = current.get(key, 0) + n
    rules = [{"src": s, "dst": d, "count": c} for (s, d), c in current.items()]
    rules.sort(key=lambda r: (-int(r["count"]), -len(str(r["src"]))))
    return rules[:400]


def append_gold_rows(rows: list[dict], filename: str, audio_rel: str | None, session_id: str) -> int:
    ensure_gold_dirs()
    path = manifest_path()
    existing: list[str] = []
    seen: set[str] = set()
    if path.is_file():
        for line in path.read_text(encoding="utf-8").splitlines():
            if not line.strip():
                continue
            try:
                old = json.loads(line)
            except json.JSONDecodeError:
                existing.append(line)
                continue
            uid = str(old.get("utterance_id") or old.get("id") or "")
            if uid:
                seen.add(uid)
            existing.append(line)
    started = time.time()
    written = 0
    for i, row in enumerate(rows):
        uid = f"{session_id}-{i:04d}"
        if uid in seen:
            continue
        rec = build_record(
            utterance_id=uid,
            session_id=session_id,
            text=str(row.get("text") or ""),
            speaker=str(row.get("speaker") or "UNKNOWN"),
            t_start=float(row.get("t_start") or 0),
            t_end=float(row.get("t_end") or 0),
            filename=filename,
            started_at=started,
            audio_rel=audio_rel,
        )
        rec["asr"] = row.get("asr") or ""
        existing.append(json.dumps(rec, ensure_ascii=False))
        written += 1
    path.write_text("\n".join(existing) + ("\n" if existing else ""), encoding="utf-8")
    return written


def corpus_stats() -> dict:
    path = manifest_path()
    counts = {"train": 0, "dev": 0, "test": 0, "total": 0, "with_audio": 0}
    if path.is_file():
        for line in path.read_text(encoding="utf-8").splitlines():
            if not line.strip():
                continue
            try:
                row = json.loads(line)
            except json.JSONDecodeError:
                continue
            split = str(row.get("split") or "train")
            counts[split] = counts.get(split, 0) + 1
            counts["total"] += 1
            if row.get("audio"):
                counts["with_audio"] += 1
    rules = load_learned()
    phrases = load_glossary_phrases()
    return {
        "ok": True,
        "manifest": str(path),
        "counts": counts,
        "rules": rules[:80],
        "rule_count": len(rules),
        "glossary_phrase_count": len(phrases),
        "excel_phrase_count": len(excel_phrases()),
        "vocabulary_revision": vocabulary_revision(),
        "recipe": recipe_status(),
    }


def _align(gold: list[dict], asr_turns: list[dict]) -> list[dict]:
    if not asr_turns:
        return gold
    used = [False] * len(asr_turns)
    for row in gold:
        t0 = float(row.get("t_start") or 0)
        best_i = None
        best_d = 9e9
        for i, turn in enumerate(asr_turns):
            if used[i]:
                continue
            d = abs(float(turn.get("t_start") or 0) - t0)
            if d < best_d:
                best_d = d
                best_i = i
        if best_i is not None and (best_d <= 8 or len(gold) == len(asr_turns)):
            used[best_i] = True
            text = str(asr_turns[best_i].get("text") or asr_turns[best_i].get("asr_text") or "")
            if not row.get("asr"):
                row["asr"] = text
        elif not row.get("asr"):
            # order fallback
            for i, turn in enumerate(asr_turns):
                if not used[i]:
                    used[i] = True
                    row["asr"] = str(turn.get("text") or turn.get("asr_text") or "")
                    break
    return gold


def start_ingest(audio: Path | None, excel: Path, audio_name: str = "", excel_name: str = "") -> str:
    job_id = uuid.uuid4().hex[:12]
    with _JOBS_LOCK:
        _JOBS[job_id] = {
            "id": job_id,
            "percent": 2.0,
            "stage": "Đã nhận file…",
            "done": False,
            "error": "",
            "preview": [],
            "learned": [],
            "gold_added": 0,
            "created": time.time(),
        }
    threading.Thread(
        target=_run_ingest,
        args=(job_id, audio, excel, audio_name, excel_name),
        daemon=True,
    ).start()
    return job_id


def job_snapshot(job_id: str) -> dict | None:
    with _JOBS_LOCK:
        job = _JOBS.get(job_id)
        if not job:
            return None
        return dict(job)


def _update(job_id: str, **fields: object) -> None:
    with _JOBS_LOCK:
        job = _JOBS.get(job_id)
        if job:
            job.update(fields)


def _run_ingest(job_id: str, audio: Path | None, excel: Path, audio_name: str, excel_name: str) -> None:
    try:
        _update(job_id, percent=8, stage="Đang đọc Excel hội thoại…")
        gold = parse_excel(excel)
        if not gold:
            raise ValueError("Excel không có dòng hội thoại (cần cột text/gold/speaker).")
        session_id = "ft-" + job_id
        ensure_gold_dirs()
        audio_rel = None
        saved_audio = None
        if audio is not None and audio.is_file():
            dest_dir = audio_root()
            dest_dir.mkdir(parents=True, exist_ok=True)
            dest = dest_dir / f"{session_id}{Path(audio_name or audio.name).suffix or audio.suffix}"
            shutil.copy2(audio, dest)
            saved_audio = dest
            audio_rel = str(dest)
        asr_turns: list[dict] = []
        if saved_audio is not None and media_transcribe is not None:
            _update(job_id, percent=20, stage="Đang ghi lời English từ file sóng…")

            def on_progress(pct: float, stage: str, text: str = "") -> None:
                _update(job_id, percent=20 + min(55, pct * 0.55), stage=stage or "Đang ghi lời…")

            _text, turns, err = media_transcribe.transcribe_with_turns(saved_audio, on_progress=on_progress)
            if err:
                _update(job_id, stage="Ghi lời máy lỗi, vẫn nạp gold Excel: " + err)
            else:
                asr_turns = turns or []
        gold = _align(gold, asr_turns)
        _update(job_id, percent=82, stage="Đang ghi corpus vàng và học cách ghi…")
        added = append_gold_rows(gold, audio_name or excel_name or "finetune", audio_rel, session_id)
        bucket: dict[tuple[str, str], int] = {}
        for row in gold:
            learn_from_pair(str(row.get("asr") or ""), str(row.get("text") or ""), bucket)
        rules = merge_learned(bucket)
        save_learned(rules)
        try:
            seed_glossary_phrases()
        except OSError:
            pass
        new_rules = [{"src": s, "dst": d, "count": n} for (s, d), n in sorted(bucket.items(), key=lambda x: -x[1])]
        preview = [
            {
                "speaker": r.get("speaker"),
                "gold": r.get("text"),
                "asr": r.get("asr") or "",
                "t_start": r.get("t_start"),
            }
            for r in gold[:40]
        ]
        _update(
            job_id,
            percent=100,
            stage="Xong. Đã cập nhật cách ghi VHF.",
            done=True,
            error="",
            preview=preview,
            learned=new_rules[:40],
            gold_added=added,
            stats=corpus_stats(),
        )
    except Exception as exc:
        _update(job_id, done=True, error=str(exc), percent=100, stage="Lỗi.")
    finally:
        for path in (audio, excel):
            if path is None:
                continue
            try:
                if str(path.resolve()).lower().startswith(str(tempfile_dir().resolve()).lower()):
                    path.unlink(missing_ok=True)
            except Exception:
                pass


def tempfile_dir() -> Path:
    return Path(tempfile.gettempdir())
