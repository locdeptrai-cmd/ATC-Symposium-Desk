"""Export VATM Excel glossary into data/vatm-terminology.tsv for build_library.py."""
from __future__ import annotations

import csv
import unicodedata
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "data" / "VATM_ATC_ATM_Terminology_Database.xlsx"
OUT = ROOT / "data" / "vatm-terminology.tsv"

DOMAIN_MAP = {
    "ATC": "atc",
    "Airspace": "atm",
    "AIS": "ais",
    "MET": "met",
    "PANS-OPS": "navigation",
    "Safety": "safety",
    "FRMS": "frms",
}

FIELDS = ("domain", "abbr", "en", "vi", "note", "source")


def normalize(s: str) -> str:
    folded = str(s or "").lower().replace("đ", "d")
    stripped = "".join(
        c for c in unicodedata.normalize("NFD", folded) if unicodedata.category(c) != "Mn"
    )
    return " ".join(stripped.split())


def clean(value) -> str:
    return " ".join(str(value or "").replace("\t", " ").split())


def note_from_master(row: dict) -> str:
    parts = [
        clean(row.get("ICAO_Ref")),
        clean(row.get("Ops_Usage")),
        clean(row.get("Definition_EN")),
    ]
    return " — ".join(p for p in parts if p)


def sheet_dicts(ws, id_header: str):
    header = None
    for row in ws.iter_rows(values_only=True):
        if row and row[0] == id_header:
            header = row
            continue
        if header and row and row[0]:
            yield dict(zip(header, row))


def rows_from_xlsx(path: Path) -> list[dict]:
    wb = load_workbook(path, data_only=True, read_only=True)
    seen_en = set()
    out: list[dict] = []

    for raw in sheet_dicts(wb["01_Master"], "Term_ID"):
        en = clean(raw.get("Term_EN"))
        vi = clean(raw.get("Term_VI"))
        domain = DOMAIN_MAP.get(clean(raw.get("Domain")))
        if not en or not vi or not domain:
            continue
        seen_en.add(normalize(en))
        out.append(
            {
                "domain": domain,
                "abbr": clean(raw.get("Abbreviation")),
                "en": en,
                "vi": vi,
                "note": note_from_master(raw),
                "source": "vatm",
            }
        )

    for raw in sheet_dicts(wb["09_Abbreviations"], "Abbr"):
        en = clean(raw.get("Expansion_EN"))
        vi = clean(raw.get("Expansion_VI"))
        domain = DOMAIN_MAP.get(clean(raw.get("Domain")))
        if not en or not vi or not domain:
            continue
        if normalize(en) in seen_en:
            continue
        seen_en.add(normalize(en))
        out.append(
            {
                "domain": domain,
                "abbr": clean(raw.get("Abbr")),
                "en": en,
                "vi": vi,
                "note": "Viết tắt VATM · " + clean(raw.get("Domain")),
                "source": "vatm",
            }
        )
    return out


def main() -> None:
    rows = rows_from_xlsx(XLSX)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    with OUT.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=FIELDS, delimiter="\t", lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)
    domains: dict[str, int] = {}
    for row in rows:
        domains[row["domain"]] = domains.get(row["domain"], 0) + 1
    print({"entries": len(rows), "domains": dict(sorted(domains.items()))})


if __name__ == "__main__":
    main()
