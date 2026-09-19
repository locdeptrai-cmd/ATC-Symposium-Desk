"""Import the Vietnam ATC airline Excel into spoken + library TSV."""
from __future__ import annotations

import csv
import shutil
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data"
XLSX_NAME = "Database_hang_hang_khong_ATC_Vietnam_2026.xlsx"
XLSX = DATA / XLSX_NAME
SPOKEN = DATA / "vn-airline-spoken.tsv"
CALLSIGNS = DATA / "vn-callsigns.tsv"
SOURCE = "airline2026"

SPOKEN_FIELDS = (
    "kind",
    "icao",
    "iata",
    "telephony",
    "spoken_en",
    "spoken_vi",
    "aliases",
    "source",
    "country",
    "confidence",
)
CALL_FIELDS = ("domain", "abbr", "en", "vi", "note", "source")


def clean(value) -> str:
    return " ".join(str(value or "").replace("\t", " ").replace("\n", " ").split())


def title_telephony(text: str) -> str:
    raw = clean(text)
    if not raw:
        return ""
    small = {"AIR", "THE", "AND", "OF"}
    parts = []
    for tok in raw.split():
        up = tok.upper()
        if up in small and parts:
            parts.append(up.title() if up != "AIR" else "Air")
            if up == "AIR":
                parts[-1] = "Air"
        else:
            parts.append(tok[:1].upper() + tok[1:].lower() if tok.isupper() or tok.islower() else tok)
    # Keep known all-caps radio names readable.
    special = {
        "VIET NAM": "Viet Nam",
        "VIETJET": "Vietjet",
        "KOREANAIR": "Korean Air",
        "JAPANAIR": "Japan Air",
        "AIRFRANS": "AirFrans",
        "TEEWAY": "TeeWay",
        "AEROHANGUK": "Aero Hanguk",
        "ALL NIPPON": "All Nippon",
        "RED CAP": "Red Cap",
        "COOL RED": "Cool Red",
        "RED NAGA": "Red Naga",
        "SUN LUX": "Sun Lux",
        "THAI ASIA": "Thai Asia",
        "THAI VIETJET": "Thai Vietjet",
    }
    key = raw.upper()
    return special.get(key, " ".join(parts))


def header_map(row: tuple) -> dict[str, int]:
    out = {}
    for i, cell in enumerate(row):
        key = clean(cell).lower()
        if key:
            out[key] = i
    return out


def cell(row: tuple, idx: dict[str, int], *names: str) -> str:
    for name in names:
        if name in idx and idx[name] < len(row):
            return clean(row[idx[name]])
    return ""


def load_old_spoken() -> tuple[dict[str, dict], list[dict]]:
    by_icao: dict[str, dict] = {}
    stations: list[dict] = []
    if not SPOKEN.is_file():
        return by_icao, stations
    with SPOKEN.open(encoding="utf-8", newline="") as fh:
        for row in csv.DictReader(fh, delimiter="\t"):
            kind = (row.get("kind") or "").strip()
            if kind == "station":
                stations.append(row)
                continue
            icao = (row.get("icao") or "").strip().upper()
            if icao:
                by_icao[icao] = row
    return by_icao, stations


def rows_from_xlsx(path: Path) -> list[dict]:
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb["AIRLINE_DB"]
    idx = None
    out: list[dict] = []
    for row in ws.iter_rows(values_only=True):
        if not row or not any(row):
            continue
        if idx is None:
            if clean(row[0]).upper() == "STT" or "icao" in clean(row[4] if len(row) > 4 else "").lower():
                idx = header_map(row)
            continue
        icao = cell(row, idx, "icao 3ld", "icao").upper()
        airline = cell(row, idx, "airline")
        telephony = cell(row, idx, "atc telephony callsign", "atc callsign")
        if not icao or len(icao) != 3 or not icao.isalpha() or not airline:
            continue
        iata = cell(row, idx, "iata")
        country = cell(row, idx, "quốc gia/vùng", "quoc gia/vung")
        kind_ops = cell(row, idx, "loại khai thác", "loai khai thac")
        status = cell(row, idx, "tình trạng/ghi chú", "tinh trang/ghi chu")
        confidence = cell(row, idx, "độ tin cậy", "do tin cay") or "High"
        example_id = cell(row, idx, "ví dụ flight id", "vi du flight id")
        example_read = cell(row, idx, "ví dụ đọc atc", "vi du doc atc")
        aliases = cell(row, idx, "alias tìm kiếm", "alias tim kiem")
        src = cell(row, idx, "nguồn tham chiếu", "nguon tham chieu")
        aliases = aliases.replace(";", ",")
        out.append(
            {
                "icao": icao,
                "iata": iata,
                "airline": airline,
                "telephony": telephony,
                "country": country,
                "kind_ops": kind_ops,
                "status": status,
                "confidence": confidence,
                "example_id": example_id,
                "example_read": example_read,
                "aliases": aliases,
                "src": src,
            }
        )
    return out


_DIGIT_TAIL = {
    "zero", "one", "two", "three", "tree", "four", "five", "fife",
    "six", "seven", "eight", "nine", "niner",
}


def merge_aliases(*chunks: str) -> str:
    seen: set[str] = set()
    out: list[str] = []
    for chunk in chunks:
        for part in str(chunk or "").replace(";", ",").split(","):
            item = clean(part)
            key = item.lower()
            if not item or key in seen:
                continue
            last = key.split()[-1]
            if last in _DIGIT_TAIL or last.isdigit():
                continue
            seen.add(key)
            out.append(item)
    return ",".join(out)


def build_spoken(airlines: list[dict], old: dict[str, dict], stations: list[dict]) -> list[dict]:
    rows: list[dict] = []
    for item in airlines:
        prev = old.get(item["icao"], {})
        spoken_en = title_telephony(item["telephony"]) or (prev.get("spoken_en") or item["airline"])
        spoken_vi = (prev.get("spoken_vi") or "").strip() or item["airline"]
        aliases = merge_aliases(
            prev.get("aliases") or "",
            item["aliases"],
            item["airline"],
            item["telephony"],
            item["iata"],
            item["icao"],
        )
        source = (
            f"{SOURCE}; rà soát 2026-09-18. {item['kind_ops']}; {item['status']}; "
            f"tin cậy {item['confidence']}. Ví dụ: {item['example_read'] or item['example_id']}. "
            f"{item['src']}".strip()
        )
        rows.append(
            {
                "kind": "airline",
                "icao": item["icao"],
                "iata": item["iata"],
                "telephony": item["telephony"].upper(),
                "spoken_en": spoken_en,
                "spoken_vi": spoken_vi,
                "aliases": aliases,
                "source": source,
                "country": item["country"],
                "confidence": item["confidence"],
            }
        )
    for st in stations:
        row = {k: st.get(k) or "" for k in SPOKEN_FIELDS}
        row["kind"] = "station"
        rows.append(row)
    return rows


def build_callsigns(airlines: list[dict]) -> list[dict]:
    rows: list[dict] = []
    for item in airlines:
        tel = item["telephony"].upper()
        note = (
            f"ICAO {item['icao']}"
            + (f". IATA {item['iata']}" if item["iata"] else "")
            + f". Telephony: {tel}. {item['country']}. {item['kind_ops']}. "
            f"{item['status']}. Độ tin cậy {item['confidence']}. "
            f"Đọc: {item['example_read'] or (tel + ' + số hiệu')}. "
            f"Đối chiếu Doc 8585 / FPL khi có xung đột."
        )
        rows.append(
            {
                "domain": "callsign",
                "abbr": item["icao"],
                "en": tel,
                "vi": item["airline"],
                "note": note,
                "source": SOURCE,
            }
        )
        if item["iata"] and item["iata"].upper() != item["icao"]:
            rows.append(
                {
                    "domain": "callsign",
                    "abbr": item["iata"],
                    "en": item["airline"],
                    "vi": f"{item['airline']} (IATA {item['iata']})",
                    "note": f"IATA {item['iata']}. Callsign vô tuyến: {tel}. ICAO {item['icao']}. {item['country']}.",
                    "source": SOURCE,
                }
            )
    return rows


def write_tsv(path: Path, fields: tuple[str, ...], rows: list[dict]) -> None:
    with path.open("w", encoding="utf-8", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=fields, delimiter="\t", extrasaction="ignore", lineterminator="\n")
        w.writeheader()
        w.writerows(rows)


def locate_xlsx() -> Path:
    if XLSX.is_file():
        return XLSX
    downloads = Path.home() / "Downloads" / XLSX_NAME
    if downloads.is_file():
        DATA.mkdir(parents=True, exist_ok=True)
        shutil.copy2(downloads, XLSX)
        return XLSX
    raise SystemExit(f"Khong thay {XLSX_NAME}")


def main() -> None:
    src = locate_xlsx()
    airlines = rows_from_xlsx(src)
    if len(airlines) < 20:
        raise SystemExit(f"Qua it hang: {len(airlines)}")
    old, stations = load_old_spoken()
    spoken = build_spoken(airlines, old, stations)
    calls = build_callsigns(airlines)
    write_tsv(SPOKEN, SPOKEN_FIELDS, spoken)
    write_tsv(CALLSIGNS, CALL_FIELDS, calls)
    vn = sum(1 for r in spoken if r["kind"] == "airline" and "việt nam" in (r.get("country") or "").lower())
    print(
        {
            "xlsx": str(src),
            "airlines": len(airlines),
            "spoken_airlines": sum(1 for r in spoken if r["kind"] == "airline"),
            "stations": sum(1 for r in spoken if r["kind"] == "station"),
            "vn_carriers": vn,
            "callsign_rows": len(calls),
        }
    )


if __name__ == "__main__":
    main()
